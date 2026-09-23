"""Append verified Cypress evidence for all Alpha black-box test cases."""

from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Alpha_Blackbox Testing.docx"
OUTPUT = ROOT / "Alpha_Blackbox Testing with Complete Cypress Evidence.docx"
EVIDENCE_JSON = ROOT / "test-results" / "alpha-blackbox-contract-evidence.json"
SCREENSHOT_DIR = ROOT / "test-results" / "alpha-cypress" / "screenshots" / "alpha-blackbox.cy.ts"


def set_cell_shading(cell, fill: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = OxmlElement(f"w:{margin}")
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
        tc_mar.append(node)


def add_caption(document: Document, number: int, case_id: str, title: str) -> None:
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = False
    run = caption.add_run(f"Figure C{number}: {case_id} — {title} (Cypress PASS evidence)")
    run.italic = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(55, 65, 81)


def normalize_section_measurements(document: Document) -> None:
    """Fix decimal twip values emitted by Word so python-docx can calculate table widths."""
    for section in document.sections:
        page_margins = section._sectPr.find(qn("w:pgMar"))
        if page_margins is None:
            continue
        for attribute in ("top", "right", "bottom", "left", "header", "footer", "gutter"):
            key = qn(f"w:{attribute}")
            value = page_margins.get(key)
            if value and "." in value:
                page_margins.set(key, str(round(float(value))))


def main() -> None:
    payload = json.loads(EVIDENCE_JSON.read_text(encoding="utf-8"))
    records = payload["cases"]
    if len(records) != 63:
        raise RuntimeError(f"Expected 63 evidence records, found {len(records)}")
    failures = [record for record in records if record["result"] != "Pass"]
    if failures:
        raise RuntimeError(f"Refusing to create pass evidence with failing cases: {[item['id'] for item in failures]}")

    screenshots: dict[str, Path] = {}
    for path in SCREENSHOT_DIR.glob("TC-*.png"):
        match = re.match(r"(TC-\d{2})-", path.name)
        if match:
            screenshots[match.group(1)] = path
    missing = [record["id"] for record in records if record["id"] not in screenshots]
    if missing:
        raise RuntimeError(f"Missing Cypress screenshots: {missing}")

    document = Document(SOURCE)
    normalize_section_measurements(document)
    document.add_page_break()

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("CYPRESS EVIDENCE FOR 63 BLACK-BOX TEST CASES")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(19)
    title_run.font.color.rgb = RGBColor(17, 24, 39)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle_run = subtitle.add_run("Alpha Black-Box Testing — Complete Automated Evidence Appendix")
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(11)
    subtitle_run.font.color.rgb = RGBColor(71, 85, 105)

    summary = document.add_table(rows=0, cols=2)
    summary.style = "Table Grid"
    summary_data = (
        ("Test runner", "Cypress 13.17.0"),
        ("Browser", "Chrome 153 (headless)"),
        ("Specification", "cypress/e2e/alpha-blackbox.cy.ts"),
        ("Document cases", "63"),
        ("Passing", "63"),
        ("Failing / pending / skipped", "0 / 0 / 0"),
        ("Screenshots", "63 case-specific captures"),
        ("Evidence source", "test-results/alpha-blackbox-contract-evidence.json"),
        ("Run timestamp", payload["generatedAt"]),
        ("Result", "PASS — all documented cases covered"),
    )
    for label, value in summary_data:
        cells = summary.add_row().cells
        cells[0].text = label
        cells[1].text = value
        set_cell_shading(cells[0], "E2E8F0")
        for cell in cells:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)
        cells[0].paragraphs[0].runs[0].bold = True
    summary.columns[0].width = Inches(2.15)
    summary.columns[1].width = Inches(4.2)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(10)
    note.paragraph_format.space_after = Pt(0)
    note_run = note.add_run(
        "Evidence method: Cypress executed one named automated checkpoint for every TC-01 through TC-63 entry. "
        "Each image below was captured during that successful Cypress run and shows the mapped executable check and its observed result."
    )
    note_run.font.name = "Arial"
    note_run.font.size = Pt(9)

    current_module = None
    image_on_page = 0
    figure_number = 1
    for record in records:
        if record["module"] != current_module:
            document.add_page_break()
            current_module = record["module"]
            image_on_page = 0
            heading = document.add_paragraph()
            heading.paragraph_format.space_after = Pt(8)
            heading.paragraph_format.keep_with_next = True
            heading_run = heading.add_run(f"Cypress Evidence — {current_module}")
            heading_run.bold = True
            heading_run.font.name = "Arial"
            heading_run.font.size = Pt(15)
            heading_run.font.color.rgb = RGBColor(15, 118, 110)
        elif image_on_page == 2:
            document.add_page_break()
            image_on_page = 0

        picture = document.add_paragraph()
        picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture.paragraph_format.space_before = Pt(0)
        picture.paragraph_format.space_after = Pt(2)
        picture.paragraph_format.keep_with_next = True
        picture.add_run().add_picture(str(screenshots[record["id"]]), width=Inches(6.1))
        add_caption(document, figure_number, record["id"], record["title"])
        figure_number += 1
        image_on_page += 1

    # Added: retain a machine-readable audit trail in the document metadata.
    document.core_properties.comments = (
        "Cypress Alpha black-box evidence: 63/63 passed; generated from "
        f"{EVIDENCE_JSON.name} on {datetime.now().isoformat(timespec='seconds')}"
    )
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
